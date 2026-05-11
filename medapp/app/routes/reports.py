from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO

from flask import Blueprint, Response, request, send_file

from .api_helpers import appointments, current_doctor_id, current_patient_id, doctors, ok, patients, payments, payment_stats, role

bp = Blueprint("reports", __name__, url_prefix="/reports")


def _date_part(value: str) -> str:
    return str(value or "")[:10]


def _month_part(value: str) -> str:
    return str(value or "")[:7] or "Sin fecha"


def _doctor_map():
    return {item.get("id"): item for item in doctors()}


def _belongs_to_doctor(item: dict, doctor_id: int | None) -> bool:
    if not doctor_id:
        return False
    return str(item.get("doctor_id") or "") == str(doctor_id)


def _belongs_to_patient(item: dict, patient_id: int | None) -> bool:
    if not patient_id:
        return False
    return str(item.get("patient_id") or "") == str(patient_id)


def _payment_belongs_to_doctor(payment: dict, appointments_by_id: dict, doctor_id: int | None) -> bool:
    if not doctor_id:
        return False
    if str(payment.get("doctor_id") or "") == str(doctor_id):
        return True
    appointment = appointments_by_id.get(payment.get("appointment_id"))
    return bool(appointment and _belongs_to_doctor(appointment, doctor_id))


def _payment_belongs_to_patient(payment: dict, appointments_by_id: dict, patient_id: int | None) -> bool:
    if not patient_id:
        return False
    if str(payment.get("patient_id") or "") == str(patient_id):
        return True
    appointment = appointments_by_id.get(payment.get("appointment_id"))
    return bool(appointment and _belongs_to_patient(appointment, patient_id))


def _visible_scope() -> dict:
    appts = appointments()
    pays = payments()
    docs = doctors()
    pats = patients()
    current_role = role()
    doctor_id = current_doctor_id() if current_role == "doctor" else None
    patient_id = current_patient_id() if current_role == "paciente" else None
    appointments_by_id = {item.get("id"): item for item in appts}

    if current_role == "doctor":
        appts = [item for item in appts if _belongs_to_doctor(item, doctor_id)]
        pays = [item for item in pays if _payment_belongs_to_doctor(item, appointments_by_id, doctor_id)]
        patient_ids = {
            item.get("patient_id")
            for item in [*appts, *pays]
            if item.get("patient_id")
        }
        docs = [item for item in docs if str(item.get("id") or "") == str(doctor_id)]
        pats = [item for item in pats if item.get("id") in patient_ids]

    if current_role == "paciente":
        appts = [item for item in appts if _belongs_to_patient(item, patient_id)]
        pays = [item for item in pays if _payment_belongs_to_patient(item, appointments_by_id, patient_id)]
        doctor_ids = {
            item.get("doctor_id")
            for item in appts
            if item.get("doctor_id")
        }
        docs = [item for item in docs if item.get("id") in doctor_ids]
        pats = [item for item in pats if str(item.get("id") or "") == str(patient_id)]

    return {
        "appointments": appts,
        "payments": pays,
        "doctors": docs,
        "patients": pats,
    }


def _filtered_appointments():
    doctor_id = request.args.get("doctor_id", type=int)
    date_from = request.args.get("date_from", "")
    date_to = request.args.get("date_to", "")
    items = _visible_scope()["appointments"]
    if doctor_id:
        items = [item for item in items if int(item.get("doctor_id") or 0) == doctor_id]
    if date_from:
        items = [item for item in items if _date_part(item.get("appointment_date")) >= date_from]
    if date_to:
        items = [item for item in items if _date_part(item.get("appointment_date")) <= date_to]
    return items


def _appointments_by_status(items):
    status_labels = {
        "pending": "Pendiente",
        "confirmed": "Confirmada",
        "cancelled": "Cancelada",
        "completed": "Atendida",
        "no_show": "No asistida",
        "rescheduled": "Reprogramada",
    }
    labels = Counter(item.get("status") or "pending" for item in items)
    return [{"status": key, "label": status_labels.get(key, key), "count": value} for key, value in labels.items()]


def _appointments_by_specialty(items):
    doctors_by_id = _doctor_map()
    labels = Counter(
        (doctors_by_id.get(item.get("doctor_id")) or {}).get("specialty") or "Sin especialidad"
        for item in items
    )
    return [{"specialty": key, "count": value} for key, value in labels.items()]


def _appointments_by_doctor(items):
    doctors_by_id = _doctor_map()
    counters = defaultdict(lambda: {"appointments": 0, "completed": 0, "cancelled": 0})
    for item in items:
        doctor = doctors_by_id.get(item.get("doctor_id")) or {}
        name = doctor.get("full_name") or "Medico no especificado"
        specialty = doctor.get("specialty") or ""
        row = counters[(item.get("doctor_id"), name, specialty)]
        row["appointments"] += 1
        if item.get("status") == "completed":
            row["completed"] += 1
        if item.get("status") == "cancelled":
            row["cancelled"] += 1
    return [
        {"doctor_id": doctor_id, "name": name, "specialty": specialty, "count": values["appointments"], **values}
        for (doctor_id, name, specialty), values in counters.items()
    ]


def _appointment_details(items):
    scope = _visible_scope()
    doctors_by_id = {item.get("id"): item for item in scope["doctors"]}
    patients_by_id = {item.get("id"): item for item in scope["patients"]}
    result = []
    for item in items:
        doctor = doctors_by_id.get(item.get("doctor_id")) or {}
        patient = patients_by_id.get(item.get("patient_id")) or {}
        result.append({
            **item,
            "doctor_name": doctor.get("full_name") or "Medico no especificado",
            "doctor_specialty": doctor.get("specialty") or "",
            "patient_name": patient.get("full_name") or "Paciente no especificado",
        })
    return sorted(result, key=lambda row: (_date_part(row.get("appointment_date")), str(row.get("appointment_time") or "")))


def _revenue_by_month(items):
    totals = defaultdict(float)
    for item in items:
        if item.get("status") != "paid":
            continue
        month = _month_part(item.get("paid_at") or item.get("created_at") or item.get("date"))
        totals[month] += float(item.get("amount") or 0)
    return [{"month": key, "amount": value} for key, value in sorted(totals.items())]


def _monthly_revenue(items):
    return [
        {"month": item["month"], "revenue": round(float(item.get("amount") or 0), 2)}
        for item in _revenue_by_month(items)[-6:]
    ]


def _payment_methods(items):
    counters = defaultdict(lambda: {"count": 0, "total": 0.0})
    for item in items:
        if item.get("status") != "paid":
            continue
        method = item.get("method") or "Otro"
        counters[method]["count"] += 1
        counters[method]["total"] += float(item.get("amount") or 0)
    return [
        {"method": method, "count": values["count"], "total": round(values["total"], 2)}
        for method, values in counters.items()
    ]


def _new_patients_by_month(items):
    totals = Counter(_month_part(item.get("created_at")) for item in items if item.get("created_at"))
    return [{"month": key, "count": value} for key, value in sorted(totals.items())[-6:]]


def _reports_payload():
    scope = _visible_scope()
    pays = scope["payments"]
    appts = _filtered_appointments()
    stats = payment_stats(pays)
    monthly = _monthly_revenue(pays)
    by_status = _appointments_by_status(appts)
    top_doctors = _appointments_by_doctor(appts)
    new_patients = _new_patients_by_month(scope["patients"])
    return {
        "summary": stats,
        "role": role(),
        "doctors": scope["doctors"],
        "revenue_by_month": _revenue_by_month(pays),
        "appointments_by_status": by_status,
        "appointments_by_specialty": _appointments_by_specialty(appts),
        "payments": pays,
        "monthly_revenue": monthly,
        "by_status": by_status,
        "top_doctors": top_doctors,
        "payment_methods": _payment_methods(pays),
        "new_patients": new_patients,
        "total_revenue": round(sum(item["revenue"] for item in monthly), 2),
        "total_appointments": len(appts),
        "active_doctors": len({item.get("doctor_id") for item in appts if item.get("doctor_id")}),
        "new_patients_total": sum(item["count"] for item in new_patients),
    }


@bp.route("/api/bootstrap")
def api_bootstrap():
    return ok(_reports_payload())


@bp.route("/api/medical-summary")
def api_medical_summary():
    appts = _filtered_appointments()
    detailed_appts = _appointment_details(appts)
    doctor_id = request.args.get("doctor_id", type=int)
    date_from = request.args.get("date_from", "")
    date_to = request.args.get("date_to", "")
    return ok({
        "doctors": _visible_scope()["doctors"],
        "filters": {
            "doctor_id": doctor_id,
            "date_from": date_from,
            "date_to": date_to,
        },
        "total_appointments": len(appts),
        "totals": {
            "appointments": len(appts),
            "completed": len([item for item in appts if item.get("status") == "completed"]),
            "cancelled": len([item for item in appts if item.get("status") == "cancelled"]),
        },
        "by_doctor": _appointments_by_doctor(appts),
        "by_status": _appointments_by_status(appts),
        "items": detailed_appts,
        "appointments": detailed_appts,
    })


@bp.route("/export/excel")
def export_excel():
    from openpyxl import Workbook

    payload = _reports_payload()
    appts = _filtered_appointments()
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["Metrica", "Valor"])
    for key, value in payload["summary"].items():
        summary.append([key, value])

    payments_sheet = workbook.create_sheet("Pagos")
    payments_sheet.append(["ID", "Paciente", "Monto", "Estado", "Metodo", "Referencia"])
    for item in payload["payments"]:
        payments_sheet.append([
            item.get("id"),
            item.get("patient_id"),
            item.get("amount"),
            item.get("status"),
            item.get("method"),
            item.get("reference"),
        ])

    appointments_sheet = workbook.create_sheet("Citas")
    appointments_sheet.append(["ID", "Fecha", "Hora", "Medico", "Paciente", "Estado", "Motivo"])
    for item in appts:
        appointments_sheet.append([
            item.get("id"),
            item.get("appointment_date"),
            item.get("appointment_time"),
            item.get("doctor_id"),
            item.get("patient_id"),
            item.get("status"),
            item.get("reason"),
        ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="medapp_reportes.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _format_money(value) -> str:
    return f"US$ {float(value or 0):,.2f}"


def _filter_summary() -> dict:
    doctor_id = request.args.get("doctor_id", type=int)
    date_from = request.args.get("date_from", "")
    date_to = request.args.get("date_to", "")
    doctor = (_doctor_map().get(doctor_id) or {}) if doctor_id else {}
    return {
        "doctor": doctor.get("full_name") if doctor_id else "Todos los medicos",
        "date_from": date_from or "Sin limite",
        "date_to": date_to or "Sin limite",
    }


class SimplePdf:
    width = 612
    height = 792
    margin = 44

    def __init__(self):
        self.pages: list[list[str]] = []
        self.ops: list[str] = []
        self.y = self.height - self.margin
        self.page()

    def page(self):
        if self.ops:
            self.pages.append(self.ops)
        self.ops = []
        self.y = self.height - self.margin

    def finish(self) -> bytes:
        if self.ops:
            self.pages.append(self.ops)
            self.ops = []
        return self._build()

    @staticmethod
    def esc(text) -> str:
        return str(text or "").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    @staticmethod
    def rgb(hex_color: str) -> tuple[float, float, float]:
        value = hex_color.lstrip("#")
        return tuple(int(value[index:index + 2], 16) / 255 for index in (0, 2, 4))

    def fill(self, x, y, w, h, color):
        r, g, b = self.rgb(color)
        self.ops.append(f"q {r:.3f} {g:.3f} {b:.3f} rg {x:.2f} {y:.2f} {w:.2f} {h:.2f} re f Q")

    def line(self, x1, y1, x2, y2, color="#E8EAF0", width=1):
        r, g, b = self.rgb(color)
        self.ops.append(f"q {r:.3f} {g:.3f} {b:.3f} RG {width:.2f} w {x1:.2f} {y1:.2f} m {x2:.2f} {y2:.2f} l S Q")

    def text(self, text, x, y, size=10, color="#2B2D42", bold=False):
        r, g, b = self.rgb(color)
        font = "F2" if bold else "F1"
        self.ops.append(
            f"BT /{font} {size:.2f} Tf {r:.3f} {g:.3f} {b:.3f} rg {x:.2f} {y:.2f} Td ({self.esc(text)}) Tj ET"
        )

    def wrap(self, text, max_chars):
        words = str(text or "").split()
        lines: list[str] = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) > max_chars and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines or [""]

    def ensure(self, needed):
        if self.y - needed < self.margin + 28:
            self.page()

    def section_title(self, title):
        self.ensure(38)
        self.text(title, self.margin, self.y, 13, "#2B2D42", True)
        self.y -= 12
        self.line(self.margin, self.y, self.width - self.margin, self.y, "#E8EAF0")
        self.y -= 18

    def _build(self) -> bytes:
        def obj(data: bytes) -> bytes:
            return data if isinstance(data, bytes) else data.encode("latin-1", "replace")

        rendered_pages = []
        page_count = len(self.pages)
        for index, ops in enumerate(self.pages, start=1):
            footer = list(ops)
            footer.append(f"BT /F1 8 Tf 0.420 0.439 0.580 rg 44 24 Td (MedApp - pagina {index} de {page_count}) Tj ET")
            rendered_pages.append("\n".join(footer).encode("latin-1", "replace"))

        objects: list[bytes] = [
            b"",  # catalog
            b"",  # pages
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
        ]
        page_refs = []
        for stream in rendered_pages:
            page_obj_num = len(objects) + 1
            content_obj_num = len(objects) + 2
            page_refs.append(f"{page_obj_num} 0 R")
            objects.append(obj(
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {self.width} {self.height}] "
                f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content_obj_num} 0 R >>"
            ))
            objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
        objects[1] = obj(f"<< /Type /Pages /Kids [{' '.join(page_refs)}] /Count {len(page_refs)} >>")

        pdf = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for number, data in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(f"{number} 0 obj\n".encode())
            pdf.extend(data)
            pdf.extend(b"\nendobj\n")
        xref = len(pdf)
        pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
        for offset in offsets[1:]:
            pdf.extend(f"{offset:010d} 00000 n \n".encode())
        pdf.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
        return bytes(pdf)


def _report_pdf_bytes(payload: dict, appts: list[dict]) -> bytes:
    filters = _filter_summary()
    details = _appointment_details(appts)
    status_rows = payload["appointments_by_status"]
    summary = payload["summary"]
    completed = len([item for item in appts if item.get("status") == "completed"])
    cancelled = len([item for item in appts if item.get("status") == "cancelled"])
    pdf = SimplePdf()

    pdf.fill(0, 704, pdf.width, 88, "#F2F7FB")
    pdf.fill(0, 704, 8, 88, "#5E60CE")
    pdf.text("MedApp", pdf.margin, 744, 22, "#5E60CE", True)
    pdf.text("Reporte medico y financiero", pdf.margin, 724, 14, "#2B2D42", True)
    pdf.text(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", 408, 748, 9, "#6B7094")
    pdf.text(f"Medico: {filters['doctor']}", pdf.margin, 690, 10, "#2B2D42", True)
    pdf.text(f"Periodo: {filters['date_from']} a {filters['date_to']}", pdf.margin, 674, 10, "#6B7094")
    pdf.y = 638

    cards = [
        ("Citas filtradas", len(appts), "#DDF2FB", "#287AA8"),
        ("Atendidas", completed, "#DDF8E6", "#24823B"),
        ("Canceladas", cancelled, "#FFD4E7", "#C0105E"),
        ("Ingresos pagados", _format_money(summary.get("total_revenue")), "#FFF0BF", "#8A6200"),
    ]
    card_w = 122
    for index, (label, value, bg, fg) in enumerate(cards):
        x = pdf.margin + index * (card_w + 10)
        pdf.fill(x, pdf.y - 60, card_w, 60, bg)
        pdf.text(label, x + 10, pdf.y - 20, 8, fg, True)
        pdf.text(value, x + 10, pdf.y - 45, 15, "#2B2D42", True)
    pdf.y -= 92

    pdf.section_title("Distribucion por estado")
    total_status = sum(int(row.get("count") or 0) for row in status_rows) or 1
    color_map = {
        "pending": "#FFF0BF",
        "confirmed": "#DDF8E6",
        "completed": "#DDF2FB",
        "cancelled": "#FFD4E7",
        "no_show": "#E9ECEF",
        "rescheduled": "#DDF2FB",
    }
    text_map = {
        "pending": "#8A6200",
        "confirmed": "#24823B",
        "completed": "#287AA8",
        "cancelled": "#C0105E",
        "no_show": "#5C6770",
        "rescheduled": "#287AA8",
    }
    for row in status_rows:
        count = int(row.get("count") or 0)
        percent = count / total_status
        pdf.text(row.get("label") or row.get("status"), pdf.margin, pdf.y, 10, "#2B2D42", True)
        pdf.text(f"{count} ({round(percent * 100)}%)", 500, pdf.y, 10, "#6B7094")
        pdf.fill(pdf.margin, pdf.y - 14, 468, 7, "#F1F3F7")
        pdf.fill(pdf.margin, pdf.y - 14, max(8, 468 * percent), 7, color_map.get(row.get("status"), "#DDF2FB"))
        pdf.text(" ", pdf.margin, pdf.y - 1, 1, text_map.get(row.get("status"), "#287AA8"))
        pdf.y -= 30

    pdf.y -= 8
    pdf.section_title("Detalle de citas")
    headers = [("Fecha", 44), ("Hora", 108), ("Paciente", 156), ("Medico", 290), ("Estado", 450)]
    row_h = 34

    def table_header():
        pdf.fill(44, pdf.y - 20, 524, 24, "#F8F9FA")
        for label, x in headers:
            pdf.text(label, x, pdf.y - 11, 8, "#6B7094", True)
        pdf.y -= 28

    table_header()
    if not details:
        pdf.text("No hay citas para los filtros seleccionados.", pdf.margin, pdf.y, 10, "#6B7094")
    for item in details:
        pdf.ensure(row_h + 36)
        if pdf.y > pdf.height - pdf.margin - 5:
            table_header()
        y_top = pdf.y
        status = item.get("status") or "pending"
        pdf.line(44, y_top - row_h + 5, 568, y_top - row_h + 5, "#E8EAF0")
        pdf.text(_date_part(item.get("appointment_date")) or "--", 44, y_top - 14, 8.5, "#2B2D42")
        pdf.text(str(item.get("appointment_time") or "--")[:5], 108, y_top - 14, 8.5, "#2B2D42")
        pdf.text((item.get("patient_name") or "--")[:24], 156, y_top - 14, 8.5, "#2B2D42")
        for line_index, line in enumerate(pdf.wrap(item.get("doctor_name") or "--", 26)[:2]):
            pdf.text(line, 290, y_top - 14 - (line_index * 10), 8.5, "#2B2D42", line_index == 0)
        pdf.fill(450, y_top - 22, 86, 16, color_map.get(status, "#E9ECEF"))
        pdf.text((next((row.get("label") for row in status_rows if row.get("status") == status), status) or "--")[:16], 458, y_top - 17, 7.5, text_map.get(status, "#5C6770"), True)
        pdf.y -= row_h

    return pdf.finish()


def _pdf_bytes(lines: list[str]) -> bytes:
    def esc(text: str) -> str:
        return str(text).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = ["BT", "/F1 12 Tf", "50 790 Td"]
    for index, line in enumerate(lines):
        if index:
            content.append("0 -18 Td")
        content.append(f"({esc(line)}) Tj")
    content.append("ET")
    stream = "\n".join(content).encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode())
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return bytes(pdf)


@bp.route("/export/pdf")
def export_pdf():
    payload = _reports_payload()
    appts = _filtered_appointments()
    return Response(
        _report_pdf_bytes(payload, appts),
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment; filename=medapp_reporte.pdf"},
    )
